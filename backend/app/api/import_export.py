from __future__ import annotations
import csv, io, json
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..database import get_db
from ..models import EquipmentAssignment, ImportBatch, User
from ..security import get_current_user, require_roles
from ..services.seed import _upsert_assignment
from ..services.status import recompute_assignment
from ..services.audit import write_audit
from ..serializers import assignment_dict

router = APIRouter(prefix='/api', tags=['Import / Export'])
REQUIRED_HEADERS = ['Car Model', 'Process', 'Transformer ID', 'PO Gun ID', 'Current Actual (Sep)', 'Current Target',
                    'Current Min', 'Current Max', 'Pressure Actual (Sep)', 'Pressure Target', 'Pressure Min', 'Pressure Max', 'Factory Zone']

def _as_float(value):
    if value is None or str(value).strip() == '': return None
    try: return float(value)
    except (ValueError, TypeError): return None

def _parse_workbook(file_bytes: bytes):
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True, keep_vba=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f'Unable to read workbook: {exc}') from exc
    if 'EQUIPMENT REGISTER' not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail='Workbook is missing the EQUIPMENT REGISTER sheet.')
    sheet = workbook['EQUIPMENT REGISTER']
    headers = [str(c.value).strip() if c.value is not None else '' for c in sheet[1]]
    index = {name: i for i, name in enumerate(headers) if name}
    missing = [h for h in REQUIRED_HEADERS if h not in index]
    if missing:
        raise HTTPException(status_code=422, detail='Missing required columns: ' + ', '.join(missing))
    rows, errors, seen_keys = [], [], set()
    for row_num, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in values): continue
        rec = {h: values[index[h]] if index[h] < len(values) else None for h in REQUIRED_HEADERS}
        model, process, transformer, gun, zone = [str(rec[k]).strip() if rec[k] is not None else '' for k in ['Car Model','Process','Transformer ID','PO Gun ID','Factory Zone']]
        if not all([model, process, transformer, gun, zone]):
            errors.append({'row': row_num, 'error': 'Missing model, process, transformer, PO Gun, or zone.'}); continue
        key = '|'.join([model, process, transformer, gun])
        if key in seen_keys:
            errors.append({'row': row_num, 'error': 'Duplicate equipment identity within uploaded file.', 'source_key': key}); continue
        seen_keys.add(key)
        rows.append({
            'source_row': row_num, 'source_key': key, 'car_model': model, 'process': process, 'transformer_id': transformer,
            'po_gun_id': gun, 'current_actual': _as_float(rec['Current Actual (Sep)']), 'current_target': _as_float(rec['Current Target']),
            'current_min': _as_float(rec['Current Min']), 'current_max': _as_float(rec['Current Max']),
            'pressure_actual': _as_float(rec['Pressure Actual (Sep)']), 'pressure_target': _as_float(rec['Pressure Target']),
            'pressure_min': _as_float(rec['Pressure Min']), 'pressure_max': _as_float(rec['Pressure Max']),
            'zone': zone, 'map_position': str(values[index['Map Position']]).strip() if 'Map Position' in index and index['Map Position'] < len(values) and values[index['Map Position']] else None,
            'snapshot_label': 'Imported workbook snapshot',
            'current_unit': 'Source unit (not labelled in register)', 'pressure_unit': 'Source unit (not labelled in register)',
        })
    workbook.close()
    return rows, errors

@router.post('/import/equipment')
async def import_equipment(file: UploadFile = File(...), confirm: bool = Form(False), db: Session = Depends(get_db),
                           user: User = Depends(require_roles('admin', 'engineer'))):
    filename = file.filename or 'upload.xlsx'
    if not filename.lower().endswith(('.xlsx', '.xlsm')):
        raise HTTPException(status_code=415, detail='Upload an .xlsx or .xlsm workbook.')
    content = await file.read()
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=413, detail='Workbook must be 25 MB or smaller.')
    rows, errors = _parse_workbook(content)
    if not rows:
        return {'preview_only': not confirm, 'source_file': filename, 'row_count': 0, 'would_create': 0, 'would_update': 0, 'errors': errors or [{'error':'No valid data rows found.'}], 'message':'No changes applied.'}
    existing = {a.source_key: a.id for a in db.scalars(select(EquipmentAssignment)).all()}
    create_count = sum(1 for x in rows if x['source_key'] not in existing)
    update_count = len(rows) - create_count
    preview = {'preview_only': not confirm, 'source_file': filename, 'row_count': len(rows), 'would_create': create_count,
               'would_update': update_count, 'rejected_count': len(errors), 'errors': errors[:100],
               'data_safety': 'No existing rows will be deleted. Source units remain unconverted.'}
    if not confirm:
        return preview
    try:
        created = updated = 0
        for record in rows:
            _, new = _upsert_assignment(db, record)
            if new: created += 1
            else: updated += 1
        batch = ImportBatch(source_file=filename, source_sheet='EQUIPMENT REGISTER', row_count=len(rows), created_count=created,
                            updated_count=updated, rejected_count=len(errors), errors_json=json.dumps(errors[:100]), imported_by_id=user.id)
        db.add(batch)
        write_audit(db, user.id, 'equipment.import', 'import_batch', None,
                    {'file': filename, 'rows': len(rows), 'created': created, 'updated': updated, 'rejected': len(errors)})
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail='Import failed. The database transaction was rolled back; no partial changes were committed.')
    return {**preview, 'preview_only': False, 'created': created, 'updated': updated, 'message': 'Import committed. Existing records were updated by identity; no records were deleted.'}

@router.get('/export/equipment.csv')
def export_equipment(model: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    stmt = select(EquipmentAssignment).order_by(EquipmentAssignment.car_model_id, EquipmentAssignment.zone_id, EquipmentAssignment.source_row)
    rows = db.scalars(stmt).all()
    if model:
        rows = [x for x in rows if x.car_model and x.car_model.name == model]
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(['Car Model','Process','Transformer ID','PO Gun ID','Current Actual (Source Unit)','Current Target','Current Min','Current Max',
                     'Pressure Actual (Source Unit)','Pressure Target','Pressure Min','Pressure Max','Factory Zone','Current Status','Pressure Status','Overall Status','Source Row'])
    for a in rows:
        writer.writerow([a.car_model.name if a.car_model else '', a.process.name if a.process else '', a.transformer.code if a.transformer else '',
                         a.po_gun.code if a.po_gun else '', a.current_actual, a.current_target, a.current_min, a.current_max,
                         a.pressure_actual, a.pressure_target, a.pressure_min, a.pressure_max, a.zone.name if a.zone else '',
                         a.current_status, a.pressure_status, a.overall_status, a.source_row])
    output.seek(0)
    filename = 'po-gun-equipment-register.csv'
    return StreamingResponse(iter([output.getvalue()]), media_type='text/csv; charset=utf-8', headers={'Content-Disposition': f'attachment; filename="{filename}"'})

@router.get('/imports')
def list_imports(db: Session = Depends(get_db), user: User = Depends(require_roles('admin', 'engineer'))):
    rows = db.query(ImportBatch).order_by(ImportBatch.imported_at.desc()).limit(100).all()
    return [{'id': x.id, 'source_file': x.source_file, 'source_sheet': x.source_sheet, 'row_count': x.row_count,
             'created_count': x.created_count, 'updated_count': x.updated_count, 'rejected_count': x.rejected_count,
             'imported_at': x.imported_at.isoformat() if x.imported_at else None} for x in rows]

@router.get('/audit')
def list_audit(offset: int = 0, limit: int = 100, db: Session = Depends(get_db), user: User = Depends(require_roles('admin'))):
    rows = db.query(ImportBatch).order_by(ImportBatch.imported_at.desc()).limit(1).all()
    from ..models import AuditLog
    logs = db.query(AuditLog).order_by(AuditLog.created_at.desc()).offset(max(0, offset)).limit(min(max(1, limit), 500)).all()
    return [{'id': x.id, 'username': x.user.username if x.user else None, 'action': x.action, 'entity_type': x.entity_type,
             'entity_id': x.entity_id, 'details': json.loads(x.details_json) if x.details_json else None,
             'created_at': x.created_at.isoformat() if x.created_at else None} for x in logs]
